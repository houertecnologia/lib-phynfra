# -*- coding: utf-8 -*-

import io
import re
import os
import docx
import pandas

from urllib.parse import urlparse
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from PIL import Image

from phynfra.aws.s3 import S3

class Word ():
	'''
	Word document inspection and dynamic report templating manager
	'''

	def __init__ (self, template, s3 = None):
		'''
		template:str|io.BytesIO|bytes - Template file path, BytesIO object or S3 URI
		s3:S3 - Optional instance of phynfra.aws.s3.S3
		'''

		self._s3 = s3

		if isinstance(template, io.BytesIO):

			template.seek(0)

			self._buffer = io.BytesIO(template.getvalue())

		elif isinstance(template, (bytes, bytearray)):

			self._buffer = io.BytesIO(template)

		elif isinstance(template, str):

			if template.startswith(('s3://', 's3a://')):

				if (not s3) or (not isinstance(s3, S3)):

					raise ValueError('An instance of S3 is required to download template from S3 URI: %s' % template)

				parsed = urlparse(template)

				bucket = parsed.netloc

				key = parsed.path.lstrip('/')

				content = s3.read(bucket = bucket, keyPrefix = key, decoder = None)

				if not content:

					raise ValueError('Failed to download template from S3: %s' % template)

				self._buffer = io.BytesIO(content)

			else:

				with open(template, 'rb') as fileHandler:

					self._buffer = io.BytesIO(fileHandler.read())

		else:

			raise ValueError('Unsupported template type: %s' % type(template))

		self._buffer.seek(0)

		self._document = Document(self._buffer)

	def placeholders (self) -> dict:
		'''
		Scans paragraphs and table cells to extract tags in r'\{\$([a-zA-Z0-9_]+)\}' format
		Returns categorized dict {'inline': [...], 'table': [...], 'image': [...]}
		'''

		pattern = re.compile(r'\{\$([a-zA-Z0-9_]+)\}')

		allTags = []

		for paragraph in self._document.paragraphs:

			matches = pattern.findall(paragraph.text)

			allTags.extend(matches)

		for table in self._document.tables:

			for row in table.rows:

				for cell in row.cells:

					for paragraph in cell.paragraphs:

						matches = pattern.findall(paragraph.text)

						allTags.extend(matches)

		inlineTags = []

		tableTags = []

		imageTags = []

		for tag in allTags:

			if tag.startswith('in_'):

				inlineTags.append(tag)

			elif tag.startswith('tb_'):

				tableTags.append(tag)

			elif tag.startswith('im_'):

				imageTags.append(tag)

		return {
			'inline': list(dict.fromkeys(inlineTags)),
			'table': list(dict.fromkeys(tableTags)),
			'image': list(dict.fromkeys(imageTags))
		}

	def render (self, context: dict) -> 'Word':
		'''
		Renders template replacing placeholders with context data
		context:dict - Mapping of tag names to replacement data
		'''

		if (not context) or (not isinstance(context, dict)):

			return self

		def replaceInParagraph (paragraph, tagPattern, replacement):

			if tagPattern in paragraph.text:

				for run in paragraph.runs:

					if tagPattern in run.text:

						run.text = run.text.replace(tagPattern, str(replacement))

				if tagPattern in paragraph.text:

					paragraph.text = paragraph.text.replace(tagPattern, str(replacement))

		for key, value in context.items():

			tagPlaceholder = '{$%s}' % key

			if key.startswith('in_'):

				for paragraph in self._document.paragraphs:

					replaceInParagraph(paragraph, tagPlaceholder, value)

				for table in self._document.tables:

					for row in table.rows:

						for cell in row.cells:

							for paragraph in cell.paragraphs:

								replaceInParagraph(paragraph, tagPlaceholder, value)

			elif key.startswith('im_'):

				for paragraph in self._document.paragraphs:

					if tagPlaceholder in paragraph.text:

						paragraph.text = paragraph.text.replace(tagPlaceholder, '').strip()

						imgBuffer = None

						if isinstance(value, (bytes, bytearray)):

							imgBuffer = io.BytesIO(value)

						elif isinstance(value, io.BytesIO):

							value.seek(0)

							imgBuffer = value

						elif isinstance(value, str):

							with open(value, 'rb') as f:

								imgBuffer = io.BytesIO(f.read())

						if imgBuffer:

							imgBuffer.seek(0)

							pilImage = Image.open(imgBuffer)

							widthPx, heightPx = pilImage.size

							imgBuffer.seek(0)

							run = paragraph.add_run()

							if widthPx > 600:

								run.add_picture(imgBuffer, width = Inches(5.5))

							else:

								run.add_picture(imgBuffer)

			elif key.startswith('tb_'):

				headers = []

				dataRows = []

				if isinstance(value, pandas.DataFrame):

					headers = [str(col) for col in value.columns]

					dataRows = value.values.tolist()

				elif isinstance(value, (list, tuple)) and (len(value) > 0) and isinstance(value[0], dict):

					headers = [str(col) for col in value[0].keys()]

					dataRows = [[row.get(h, '') for h in headers] for row in value]

				if len(headers) > 0:

					for paragraph in self._document.paragraphs:

						if tagPlaceholder in paragraph.text:

							paragraph.text = paragraph.text.replace(tagPlaceholder, '').strip()

							docTable = self._document.add_table(rows = len(dataRows) + 1, cols = len(headers))

							docTable.style = 'Table Grid'

							for colIdx, header in enumerate(headers):

								cell = docTable.cell(0, colIdx)

								cell.text = header

								for p in cell.paragraphs:

									for run in p.runs:

										run.bold = True

							for rowIdx, rowData in enumerate(dataRows):

								for colIdx, cellValue in enumerate(rowData):

									cell = docTable.cell(rowIdx + 1, colIdx)

									cell.text = '' if cellValue is None else str(cellValue)

							paragraph._p.addnext(docTable._tbl)

		return self

	def save (self, destination: str):
		'''
		Saves the rendered document to local path or AWS S3 URI
		destination:str - Local destination path or S3 URI (s3:// or s3a://)
		'''

		if (not destination) or (not isinstance(destination, str)) or (not destination.strip()):

			raise ValueError('Destination must be a non-empty string!')

		targetDest = destination.strip()

		if targetDest.startswith(('s3://', 's3a://')):

			if (not self._s3) or (not isinstance(self._s3, S3)):

				raise ValueError('An instance of S3 is required to save to S3 destination: %s' % targetDest)

			parsed = urlparse(targetDest)

			bucket = parsed.netloc

			key = parsed.path.lstrip('/')

			outputBuffer = io.BytesIO()

			self._document.save(outputBuffer)

			outputBuffer.seek(0)

			payloadBytes = outputBuffer.getvalue()

			return self._s3.write(
				bucket = bucket,
				payload = payloadBytes,
				keyPrefix = key,
				contentType = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
			)

		else:

			targetPath = os.path.abspath(targetDest)

			targetDir = os.path.dirname(targetPath)

			if targetDir and (not os.path.exists(targetDir)):

				os.makedirs(targetDir, exist_ok = True)

			self._document.save(targetPath)

			return targetPath
